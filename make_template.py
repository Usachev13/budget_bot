"""Собирает шаблон Google Таблицы: категории, правила, счета, лимиты."""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
import db

HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="2A78D6")
WALLET = {"salary": "зарплата", "ip": "ИП", "income": "доход", "transfer": "перевод"}
KIND = {"expense": "расход", "income": "доход", "transfer": "перевод"}


def fill_sheet(ws, headers, widths, rows):
    ws.append(headers)
    for i, cell in enumerate(ws[1], start=1):
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[chr(64 + i)].width = widths[i - 1]
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"


def build(path):
    db.init()
    with db.conn() as c:
        cats = c.execute("SELECT name,wallet,kind,monthly_limit FROM categories ORDER BY sort").fetchall()
        rules = c.execute("SELECT pattern,category FROM rules ORDER BY category,pattern").fetchall()

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Категории"
    fill_sheet(ws, ["Категория", "Кошелёк", "Тип", "Лимит ₽/мес", "Комментарий"],
               [30, 14, 12, 14, 44],
               [[r["name"], WALLET[r["wallet"]], KIND[r["kind"]],
                 r["monthly_limit"] or "", ""] for r in cats])
    dv = DataValidation(type="list", formula1='"зарплата,ИП,доход,перевод"', allow_blank=True)
    ws.add_data_validation(dv); dv.add("B2:B200")
    dv2 = DataValidation(type="list", formula1='"расход,доход,перевод"', allow_blank=True)
    ws.add_data_validation(dv2); dv2.add("C2:C200")

    ws2 = wb.create_sheet("Правила")
    fill_sheet(ws2, ["Текст в описании операции", "Категория", "Комментарий"],
               [42, 30, 40], [[r["pattern"], r["category"], ""] for r in rules])

    ws3 = wb.create_sheet("Счета")
    fill_sheet(ws3, ["Название", "Тип", "Валюта", "Назначение",
                     "Остаток на старте", "Лимит по карте"],
               [30, 16, 10, 22, 18, 16],
               [["Яндекс · Сейв 994253", "накопительный", "RUB", "свободные", 24222, ""],
                ["Яндекс · Сейв 790267", "накопительный", "RUB", "билеты", 54418, ""],
                ["Яндекс · Сейв 041732", "накопительный", "RUB", "налоги", 114935, ""],
                ["Яндекс · счёт 423455", "текущий", "RUB", "повседневный", 490, ""],
                ["Озон · карта 175228", "карта", "RUB", "", -20152, 49000],
                ["Т-Банк", "карта", "RUB", "", "", 385000],
                ["Сбер", "карта", "RUB", "", "", 100000],
                ["ОТП", "карта", "RUB", "", "", 150000],
                ["Шекели", "валютный", "ILS", "переезд", "", ""],
                ["USDT", "валютный", "USDT", "переезд", "", ""]])
    dv3 = DataValidation(type="list", formula1='"текущий,накопительный,карта,валютный"', allow_blank=True)
    ws3.add_data_validation(dv3); dv3.add("B2:B100")

    ws4 = wb.create_sheet("Как пользоваться")
    ws4.column_dimensions["A"].width = 110
    for line in [
        "Эта таблица — источник правды для бота. После правок выполни в боте команду /синхрон.",
        "",
        "Лист «Категории». Кошелёк: зарплата (лимиты в сумме = 175 000 ₽), ИП (без лимита), доход, перевод.",
        "Тип определяет, как операция попадает в отчёты: расход, доход или перевод между своими счетами.",
        "Чтобы убрать категорию, удали строку. Чтобы добавить — допиши строку снизу.",
        "",
        "Лист «Правила». Если текст из первой колонки встречается в описании операции, бот ставит эту категорию.",
        "Пишется строчными буквами, часть слова тоже подходит: «пятёроч» поймает «Пятёрочка» и «Пятёрочка Доставка».",
        "",
        "Лист «Счета». Тип: текущий, накопительный, карта, валютный. Для карт остаток пишется со знаком минус.",
        "",
        "Категории и правила из таблицы полностью заменяют то, что было в боте.",
    ]:
        ws4.append([line])

    wb.save(path)
    return path


if __name__ == "__main__":
    print(build(sys.argv[1] if len(sys.argv) > 1 else "budget_template.xlsx"))
